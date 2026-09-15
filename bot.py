import asyncio
import logging
import os
import re
import sqlite3
from datetime import datetime
from io import BytesIO

from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command, CommandStart
from aiogram.types import Message, BufferedInputFile, CallbackQuery
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# НАСТРОЙКИ
# ---------------------------------------------------------------------------

BOT_TOKEN = os.getenv("BOT_TOKEN", "ВАШ_ТОКЕН_СЮДА")
DB_PATH = os.getenv("DB_PATH", "swim.db")

# Telegram numeric ID (не username!) владельца и тренера — только у них есть
# право добавлять/менять результаты и смотреть общую таблицу.
# Узнать свой ID можно написав боту @userinfobot.
ADMIN_IDS = {
    111111111,  # <- впиши сюда свой Telegram ID
    222222222,  # <- впиши сюда Telegram ID тренера
}

logging.basicConfig(level=logging.INFO)

STROKE_ALIASES = {
    "кроль": "кроль", "вольный": "кроль", "вс": "кроль", "вольный стиль": "кроль",
    "брасс": "брасс",
    "спина": "спина", "на спине": "спина",
    "батт": "баттерфляй", "баттерфляй": "баттерфляй", "дельфин": "баттерфляй",
    "компл": "комплекс", "комплекс": "комплекс", "к/п": "комплекс",
}


def is_admin(telegram_id: int) -> bool:
    return telegram_id in ADMIN_IDS

# ---------------------------------------------------------------------------
# БАЗА ДАННЫХ
# ---------------------------------------------------------------------------

def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            distance INTEGER NOT NULL,
            stroke TEXT NOT NULL,
            time_seconds REAL NOT NULL,
            session_type TEXT NOT NULL DEFAULT 'тренировка',
            recorded_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            telegram_id INTEGER PRIMARY KEY,
            name TEXT NOT NULL
        )
        """
    )
    conn.commit()
    conn.close()


def link_user(telegram_id: int, name: str):
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        "INSERT INTO users (telegram_id, name) VALUES (?, ?) "
        "ON CONFLICT(telegram_id) DO UPDATE SET name=excluded.name",
        (telegram_id, name),
    )
    conn.commit()
    conn.close()


def get_linked_name(telegram_id: int):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute("SELECT name FROM users WHERE telegram_id = ?", (telegram_id,))
    row = cur.fetchone()
    conn.close()
    return row[0] if row else None


def get_all_linked_users():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute("SELECT telegram_id, name FROM users ORDER BY name")
    rows = cur.fetchall()
    conn.close()
    return rows


def unlink_user(telegram_id: int) -> bool:
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute("DELETE FROM users WHERE telegram_id = ?", (telegram_id,))
    conn.commit()
    changed = cur.rowcount > 0
    conn.close()
    return changed


def save_result(name, distance, stroke, time_seconds, session_type):
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        "INSERT INTO results (name, distance, stroke, time_seconds, session_type, recorded_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (name, distance, stroke, time_seconds, session_type, datetime.now().isoformat()),
    )
    conn.commit()
    conn.close()


def delete_result(result_id: int) -> bool:
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute("DELETE FROM results WHERE id = ?", (result_id,))
    conn.commit()
    changed = cur.rowcount > 0
    conn.close()
    return changed


def update_result(result_id, name, distance, stroke, time_seconds, session_type) -> bool:
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "UPDATE results SET name=?, distance=?, stroke=?, time_seconds=?, session_type=? WHERE id=?",
        (name, distance, stroke, time_seconds, session_type, result_id),
    )
    conn.commit()
    changed = cur.rowcount > 0
    conn.close()
    return changed


def get_previous_result(name, distance, stroke):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "SELECT time_seconds, recorded_at, session_type FROM results "
        "WHERE name = ? AND distance = ? AND stroke = ? "
        "ORDER BY recorded_at DESC LIMIT 2",
        (name, distance, stroke),
    )
    rows = cur.fetchall()
    conn.close()
    if len(rows) >= 2:
        return rows[1]
    return None


def get_personal_best(name, distance, stroke):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "SELECT MIN(time_seconds) FROM results WHERE name = ? AND distance = ? AND stroke = ?",
        (name, distance, stroke),
    )
    row = cur.fetchone()
    conn.close()
    return row[0] if row else None


def get_distances_for_name(name):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "SELECT DISTINCT distance, stroke FROM results WHERE name = ? ORDER BY distance, stroke",
        (name,),
    )
    rows = cur.fetchall()
    conn.close()
    return rows


def get_entries(name, distance, stroke, session_type, limit=2):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "SELECT time_seconds, recorded_at FROM results "
        "WHERE name = ? AND distance = ? AND stroke = ? AND session_type = ? "
        "ORDER BY recorded_at DESC LIMIT ?",
        (name, distance, stroke, session_type, limit),
    )
    rows = cur.fetchall()
    conn.close()
    return rows


def get_best(name, distance, stroke, session_type):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "SELECT MIN(time_seconds) FROM results "
        "WHERE name = ? AND distance = ? AND stroke = ? AND session_type = ?",
        (name, distance, stroke, session_type),
    )
    row = cur.fetchone()
    conn.close()
    return row[0] if row else None


def get_all_results(name=None):
    conn = sqlite3.connect(DB_PATH)
    if name:
        cur = conn.execute(
            "SELECT name, distance, stroke, time_seconds, session_type, recorded_at "
            "FROM results WHERE name = ? ORDER BY name, distance, stroke, recorded_at",
            (name,),
        )
    else:
        cur = conn.execute(
            "SELECT name, distance, stroke, time_seconds, session_type, recorded_at "
            "FROM results ORDER BY name, distance, stroke, recorded_at"
        )
    rows = cur.fetchall()
    conn.close()
    return rows


def get_recent(limit=10):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        "SELECT id, name, distance, stroke, time_seconds, session_type, recorded_at "
        "FROM results ORDER BY id DESC LIMIT ?",
        (limit,),
    )
    rows = cur.fetchall()
    conn.close()
    return rows

# ---------------------------------------------------------------------------
# ПАРСИНГ
# ---------------------------------------------------------------------------

def parse_time(raw: str) -> float:
    raw = raw.strip().replace(",", ".")
    if "." in raw:
        main, frac = raw.rsplit(".", 1)
        frac_seconds = float("0." + frac)
    else:
        main = raw
        frac_seconds = 0.0
    parts = [p for p in main.split(":") if p != ""]
    if len(parts) == 1:
        total = float(parts[0])
    elif len(parts) == 2:
        m, s = parts
        total = int(m) * 60 + float(s)
    elif len(parts) == 3:
        a, b, c = parts
        if int(c) < 100 and "." not in raw:
            total = int(a) * 60 + int(b) + int(c) / 100
        else:
            total = int(a) * 3600 + int(b) * 60 + float(c)
    else:
        raise ValueError(f"не смог разобрать время: {raw}")
    return total + frac_seconds


def format_time(seconds: float) -> str:
    m = int(seconds // 60)
    s = seconds - m * 60
    if m > 0:
        return f"{m}:{s:05.2f}"
    return f"{s:.2f}"


def normalize_stroke(raw: str) -> str:
    key = raw.strip().lower()
    return STROKE_ALIASES.get(key, raw.strip())


def parse_record(text: str):
    """Формат: Имя - Дистанция - Стиль - Время [- трен/сорев]"""
    cleaned = text.replace("—", "-").replace("–", "-")
    parts = [p.strip() for p in cleaned.split("-")]
    parts = [p for p in parts if p != ""]
    if len(parts) < 4:
        return None
    name = parts[0].strip().title()
    distance_digits = re.sub(r"[^\d]", "", parts[1])
    if not distance_digits:
        return None
    distance = int(distance_digits)
    stroke = normalize_stroke(parts[2])
    try:
        time_seconds = parse_time(parts[3])
    except ValueError:
        return None
    session_type = "тренировка"
    if len(parts) >= 5:
        t = parts[4].lower()
        if "сорев" in t:
            session_type = "соревнования"
        elif "трен" in t:
            session_type = "тренировка"
    return name, distance, stroke, time_seconds, session_type

# ---------------------------------------------------------------------------
# БОТ
# ---------------------------------------------------------------------------

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

ADMIN_HELP = (
    "🏊 <b>Режим тренера/владельца</b>\n\n"
    "<b>Добавить результат</b> — сообщение в формате:\n"
    "<code>Имя - Дистанция - Стиль - Время [- трен/сорев]</code>\n"
    "Например: <code>Аня - 200 - брасс - 4:00.00 - сорев</code>\n"
    "В группе, если бот видит только команды, используй:\n"
    "<code>/add Аня - 200 - брасс - 4:00.00</code>\n\n"
    "<b>Команды</b>:\n"
    "/last — последние 10 записей с их id\n"
    "/delete id — удалить запись по id\n"
    "/edit id - Имя - Дистанция - Стиль - Время - Тип — исправить запись\n"
    "/table — Excel со всеми результатами всех пловцов\n"
    "/table Имя — Excel по одному пловцу\n"
    "/link id Имя — привязать чужой Telegram ID к имени пловца (чтобы он видел "
    "свой прогресс через /my без самостоятельной регистрации)\n"
    "/unlink id — отвязать\n"
    "/users — список всех привязанных пловцов\n"
    "/my — посмотреть свой личный прогресс (как у обычных пользователей)\n"
)

USER_HELP = (
    "🏊 Привет! Здесь можно смотреть только свои результаты.\n\n"
    "Сначала привяжи себя к имени, под которым тренер вносит твои результаты:\n"
    "<code>/iam Твоё Имя</code>\n\n"
    "После этого команда /my покажет кнопки с твоими дистанциями — жми на "
    "нужную, и бот пришлёт лучший результат на соревнованиях, лучший на "
    "тренировках и два последних результата в каждой из категорий."
)


@dp.message(CommandStart())
async def cmd_start(message: Message):
    if is_admin(message.from_user.id):
        await message.answer(ADMIN_HELP, parse_mode="HTML")
    else:
        await message.answer(USER_HELP, parse_mode="HTML")


@dp.message(Command("help"))
async def cmd_help(message: Message):
    await cmd_start(message)


@dp.message(Command("whoami"))
async def cmd_whoami(message: Message):
    uid = message.from_user.id
    name = get_linked_name(uid)
    role = "тренер/владелец" if is_admin(uid) else "пловец"
    text = f"Твой Telegram ID: <code>{uid}</code>\nРоль: {role}\n"
    text += f"Привязанное имя: {name}" if name else "Имя не привязано (используй /iam Имя)"
    await message.answer(text, parse_mode="HTML")


@dp.message(Command("iam"))
async def cmd_iam(message: Message):
    arg = message.text.partition(" ")[2].strip()
    if not arg:
        await message.answer("Напиши так: /iam Аня")
        return
    name = arg.title()
    link_user(message.from_user.id, name)
    await message.answer(f"Готово! Теперь ты — {name}. Жми /my, чтобы увидеть свой прогресс.")


@dp.message(Command("link"))
async def cmd_link(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Эта команда только для тренера/владельца.")
        return
    arg = message.text.partition(" ")[2].strip()
    telegram_id_str, _, name = arg.partition(" ")
    if not telegram_id_str.isdigit() or not name.strip():
        await message.answer(
            "Напиши так: /link 123456789 Аня\n"
            "(узнать чужой Telegram ID: пусть человек напишет боту @userinfobot "
            "или пришлёт тебе свой /whoami)"
        )
        return
    name = name.strip().title()
    link_user(int(telegram_id_str), name)
    await message.answer(f"Готово! ID {telegram_id_str} привязан к имени «{name}».")


@dp.message(Command("unlink"))
async def cmd_unlink(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Эта команда только для тренера/владельца.")
        return
    arg = message.text.partition(" ")[2].strip()
    if not arg.isdigit():
        await message.answer("Напиши так: /unlink 123456789")
        return
    ok = unlink_user(int(arg))
    await message.answer("Отвязано ✅" if ok else "Такой ID не привязан ни к кому.")


@dp.message(Command("users"))
async def cmd_users(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Эта команда только для тренера/владельца.")
        return
    rows = get_all_linked_users()
    if not rows:
        await message.answer("Пока никто не привязан. Используй /link id Имя.")
        return
    lines = ["<b>Привязанные пользователи</b>:"]
    for telegram_id, name in rows:
        lines.append(f"• {name} — <code>{telegram_id}</code>")
    await message.answer("\n".join(lines), parse_mode="HTML")

# ---------------------------------------------------------------------------
# ДОБАВЛЕНИЕ / ИЗМЕНЕНИЕ (только тренер/владелец)
# ---------------------------------------------------------------------------

async def handle_new_record(message: Message, raw_text: str):
    parsed = parse_record(raw_text)
    if not parsed:
        await message.answer(
            "Не смог разобрать запись 🤔\nФормат: <code>Имя - Дистанция - Стиль - Время</code>",
            parse_mode="HTML",
        )
        return
    name, distance, stroke, time_seconds, session_type = parsed

    prev = get_previous_result(name, distance, stroke)
    best_before = get_personal_best(name, distance, stroke)

    save_result(name, distance, stroke, time_seconds, session_type)

    lines = [
        f"✅ Записано: <b>{name}</b> — {distance}м {stroke}, "
        f"{format_time(time_seconds)} ({session_type})"
    ]
    if prev:
        delta = time_seconds - prev[0]
        if delta < 0:
            lines.append(f"⬇️ Быстрее предыдущего на {format_time(abs(delta))}")
        elif delta > 0:
            lines.append(f"⬆️ Медленнее предыдущего на {format_time(abs(delta))}")
        else:
            lines.append("➖ Точно как в прошлый раз")
    else:
        lines.append("📌 Первый результат на этой дистанции этим стилем")

    if best_before is not None:
        if time_seconds < best_before:
            lines.append("🏆 Новый личный рекорд!")
        else:
            gap = time_seconds - best_before
            lines.append(f"Личный рекорд: {format_time(best_before)} (отставание {format_time(gap)})")

    await message.answer("\n".join(lines), parse_mode="HTML")


@dp.message(Command("add"))
async def cmd_add(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Добавлять результаты может только тренер или владелец.")
        return
    text_after_command = message.text.partition(" ")[2]
    if not text_after_command.strip():
        await message.answer("Напиши так: /add Аня - 200 - брасс - 4:00.00")
        return
    await handle_new_record(message, text_after_command)


@dp.message(Command("last"))
async def cmd_last(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Эта команда только для тренера/владельца.")
        return
    rows = get_recent(10)
    if not rows:
        await message.answer("Записей пока нет.")
        return
    lines = ["<b>Последние записи</b>:"]
    for rid, name, distance, stroke, time_seconds, session_type, recorded_at in rows:
        date_str = recorded_at.split("T")[0]
        lines.append(
            f"#{rid} — {name}, {distance}м {stroke}, {format_time(time_seconds)}, "
            f"{session_type}, {date_str}"
        )
    await message.answer("\n".join(lines), parse_mode="HTML")


@dp.message(Command("delete"))
async def cmd_delete(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Эта команда только для тренера/владельца.")
        return
    arg = message.text.partition(" ")[2].strip()
    if not arg.isdigit():
        await message.answer("Напиши так: /delete 15 (посмотреть id можно через /last)")
        return
    ok = delete_result(int(arg))
    await message.answer("Удалено ✅" if ok else "Запись с таким id не найдена.")


@dp.message(Command("edit"))
async def cmd_edit(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Эта команда только для тренера/владельца.")
        return
    rest = message.text.partition(" ")[2].strip()
    id_part, _, record_part = rest.partition(" ")
    if not id_part.isdigit() or not record_part.strip():
        await message.answer(
            "Напиши так: /edit 15 - Аня - 200 - брасс - 3:58.00 - сорев\n"
            "(id смотри через /last)"
        )
        return
    parsed = parse_record(record_part)
    if not parsed:
        await message.answer("Не смог разобрать новую запись, проверь формат.")
        return
    name, distance, stroke, time_seconds, session_type = parsed
    ok = update_result(int(id_part), name, distance, stroke, time_seconds, session_type)
    if ok:
        await message.answer(
            f"Исправлено ✅ #{id_part}: {name} — {distance}м {stroke}, "
            f"{format_time(time_seconds)} ({session_type})"
        )
    else:
        await message.answer("Запись с таким id не найдена.")


def build_excel(rows, title="Результаты"):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    headers = ["Пловец", "Дистанция, м", "Стиль", "Время", "Тип", "Дата"]
    ws.append(headers)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for name, distance, stroke, time_seconds, session_type, recorded_at in rows:
        date_str = recorded_at.split("T")[0]
        ws.append([name, distance, stroke, format_time(time_seconds), session_type, date_str])

    for col_idx, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))] + [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@dp.message(Command("table"))
async def cmd_table(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Общая таблица доступна только тренеру/владельцу. Свой прогресс смотри через /my.")
        return
    arg = message.text.partition(" ")[2].strip()
    name = arg.title() if arg else None
    rows = get_all_results(name)
    if not rows:
        await message.answer("Пока нет данных для таблицы.")
        return
    buf = build_excel(rows, title=name or "Все пловцы")
    filename = f"swim_results_{name or 'all'}.xlsx"
    await message.answer_document(BufferedInputFile(buf.read(), filename=filename))

# ---------------------------------------------------------------------------
# ПРОСМОТР СВОЕГО ПРОГРЕССА (доступно всем)
# ---------------------------------------------------------------------------

@dp.message(Command("my"))
async def cmd_my(message: Message):
    uid = message.from_user.id
    name = get_linked_name(uid)
    if not name:
        await message.answer("Сначала привяжи имя: <code>/iam Твоё Имя</code>", parse_mode="HTML")
        return
    combos = get_distances_for_name(name)
    if not combos:
        await message.answer(f"Пока нет ни одного результата для {name}.")
        return
    builder = InlineKeyboardBuilder()
    for distance, stroke in combos:
        label = f"{distance}м {stroke}"
        builder.button(text=label, callback_data=f"prog:{distance}:{stroke}:{uid}")
    builder.adjust(2)
    await message.answer(f"Выбери дистанцию, {name}:", reply_markup=builder.as_markup())


@dp.callback_query(F.data.startswith("prog:"))
async def on_progress_click(callback: CallbackQuery):
    _, distance_str, stroke, owner_id_str = callback.data.split(":", 3)
    distance = int(distance_str)
    owner_id = int(owner_id_str)

    if callback.from_user.id != owner_id:
        await callback.answer("Это не твои кнопки 🙂", show_alert=True)
        return

    name = get_linked_name(owner_id)
    if not name:
        await callback.answer("Имя не найдено.", show_alert=True)
        return

    best_comp = get_best(name, distance, stroke, "соревнования")
    best_train = get_best(name, distance, stroke, "тренировка")
    last_comp = get_entries(name, distance, stroke, "соревнования", limit=2)
    last_train = get_entries(name, distance, stroke, "тренировка", limit=2)

    lines = [f"📊 <b>{name} — {distance}м {stroke}</b>\n"]

    lines.append(f"🏆 Лучший на соревнованиях: {format_time(best_comp) if best_comp else '—'}")
    lines.append(f"🏋️ Лучший на тренировках: {format_time(best_train) if best_train else '—'}\n")

    lines.append("<b>Соревнования (последние):</b>")
    if last_comp:
        for t, dt in last_comp:
            lines.append(f"• {format_time(t)} ({dt.split('T')[0]})")
        if len(last_comp) == 2:
            delta = last_comp[0][0] - last_comp[1][0]
            arrow = "⬇️ быстрее" if delta < 0 else ("⬆️ медленнее" if delta > 0 else "➖ так же")
            lines.append(f"Изменение: {arrow} на {format_time(abs(delta))}")
    else:
        lines.append("• пока нет результатов")

    lines.append("\n<b>Тренировки (последние):</b>")
    if last_train:
        for t, dt in last_train:
            lines.append(f"• {format_time(t)} ({dt.split('T')[0]})")
        if len(last_train) == 2:
            delta = last_train[0][0] - last_train[1][0]
            arrow = "⬇️ быстрее" if delta < 0 else ("⬆️ медленнее" if delta > 0 else "➖ так же")
            lines.append(f"Изменение: {arrow} на {format_time(abs(delta))}")
    else:
        lines.append("• пока нет результатов")

    await callback.message.answer("\n".join(lines), parse_mode="HTML")
    await callback.answer()

# ---------------------------------------------------------------------------
# АВТОРАСПОЗНАВАНИЕ СООБЩЕНИЙ (только тренер/владелец)
# ---------------------------------------------------------------------------

@dp.message(F.text)
async def catch_all(message: Message):
    text = message.text or ""
    if text.startswith("/"):
        return
    dash_count = text.count("-") + text.count("—") + text.count("–")
    if dash_count < 3:
        return
    if not is_admin(message.from_user.id):
        return  # обычные пользователи не могут добавлять записи
    await handle_new_record(message, text)


async def main():
    init_db()
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
